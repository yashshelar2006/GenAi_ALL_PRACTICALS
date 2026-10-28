import os
import re
import uuid
from pathlib import Path

import pymupdf
import requests
from docx import Document as DocxDocument
from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request
from werkzeug.utils import secure_filename


BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")
UPLOAD_DIR = BASE_DIR / "uploads"
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_DOCUMENTS = 10
CHUNK_SIZE = int(os.getenv("RAG_CHUNK_SIZE", "1000"))
CHUNK_OVERLAP = int(os.getenv("RAG_CHUNK_OVERLAP", "150"))
SIMILARITY_THRESHOLD = float(os.getenv("RAG_SIMILARITY_THRESHOLD", "0.30"))
EMBEDDING_MODEL_NAME = os.getenv(
	"EMBEDDING_MODEL_NAME", "sentence-transformers/all-MiniLM-L6-v2"
)

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES + 1024 * 1024

documents = {}
embedding_model = None


def get_colab_config():
	base_url = os.getenv("COLAB_API_URL", "").strip().rstrip("/")
	token = os.getenv("COLAB_API_TOKEN", "").strip()
	if not base_url or not token:
		return None, None
	return base_url, token


def get_embedding_model():
	global embedding_model
	if embedding_model is None:
		from sentence_transformers import SentenceTransformer

		embedding_model = SentenceTransformer(EMBEDDING_MODEL_NAME)
	return embedding_model


def split_into_chunks(text):
	text = re.sub(r"\s+", " ", text).strip()
	if not text:
		return []
	chunks = []
	start = 0
	while start < len(text):
		end = min(start + CHUNK_SIZE, len(text))
		if end < len(text):
			boundary = text.rfind(" ", start, end)
			if boundary > start + CHUNK_SIZE // 2:
				end = boundary
		chunks.append(text[start:end].strip())
		if end >= len(text):
			break
		start = max(end - CHUNK_OVERLAP, start + 1)
	return [chunk for chunk in chunks if chunk]


def extract_text(path, extension):
	if extension == "pdf":
		with pymupdf.open(path) as pdf:
			return "\n".join(page.get_text() for page in pdf)
	if extension == "docx":
		document = DocxDocument(path)
		paragraphs = [paragraph.text for paragraph in document.paragraphs]
		for table in document.tables:
			paragraphs.extend(
				cell.text for row in table.rows for cell in row.cells
			)
		return "\n".join(paragraphs)
	return path.read_bytes().decode("utf-8-sig")


def validate_history(value):
	if not isinstance(value, list):
		raise ValueError("Conversation history must be a list.")
	history = []
	for item in value[-12:]:
		if not isinstance(item, dict) or item.get("role") not in {"user", "assistant"}:
			continue
		content = item.get("content")
		if isinstance(content, str) and content.strip():
			history.append({"role": item["role"], "content": content.strip()[:4000]})
	return history


def retrieve_context(question, document_ids):
	if not document_ids:
		return [], []

	missing = [document_id for document_id in document_ids if document_id not in documents]
	if missing:
		raise ValueError("One or more selected documents are no longer available. Refresh the page and try again.")

	model = get_embedding_model()
	question_embedding = model.encode(question, normalize_embeddings=True)
	matches = []
	for document_id in document_ids:
		document = documents[document_id]
		scores = document["embeddings"] @ question_embedding
		for index, score in enumerate(scores):
			if float(score) >= SIMILARITY_THRESHOLD:
				matches.append((float(score), document, document["chunks"][index]))

	matches.sort(key=lambda match: match[0], reverse=True)
	selected = matches[:4]
	context = [
		f"[Source: {document['name']} | relevance: {score:.2f}]\n{chunk}"
		for score, document, chunk in selected
	]
	sources = list(dict.fromkeys(document["name"] for _, document, _ in selected))
	return context, sources


@app.get("/")
def index():
	return render_template("index.html")


@app.get("/api/health")
def health():
	base_url, token = get_colab_config()
	if not base_url or not token:
		return jsonify(
			status="degraded",
			configured=False,
			colab="not_configured",
			message="Set COLAB_API_URL and COLAB_API_TOKEN in the local .env file.",
		), 503
	try:
		response = requests.get(
			f"{base_url}/health",
			headers={"Authorization": f"Bearer {token}"},
			timeout=10,
		)
		if not response.ok:
			raise requests.RequestException("Colab health check returned an error.")
	except requests.RequestException:
		return jsonify(
			status="degraded",
			configured=True,
			colab="unavailable",
			message="The Colab API could not be reached. Check that the runtime and tunnel are active.",
		), 502
	return jsonify(status="ok", configured=True, colab="connected")


@app.post("/api/chat")
def chat():
	payload = request.get_json(silent=True)
	if not isinstance(payload, dict):
		return jsonify(error="Send a JSON object with a question."), 400

	question = payload.get("question")
	if not isinstance(question, str) or not question.strip():
		return jsonify(error="Enter a question before sending."), 400
	question = question.strip()
	if len(question) > 5000:
		return jsonify(error="Questions must be 5,000 characters or fewer."), 400

	preferences = payload.get("preferences", {})
	if not isinstance(preferences, dict):
		return jsonify(error="Preferences must be an object."), 400
	conversation_id = payload.get("conversation_id")
	if conversation_id is not None and (not isinstance(conversation_id, str) or len(conversation_id) > 100):
		return jsonify(error="Conversation identifier must be a string of 100 characters or fewer."), 400
	preferred_name = str(preferences.get("name", "")).strip()[:80]
	response_style = str(preferences.get("style", "")).strip()[:200]
	learning_goals = str(preferences.get("goals", "")).strip()[:500]
	try:
		history = validate_history(payload.get("history", []))
		document_ids = payload.get("document_ids", [])
		if not isinstance(document_ids, list) or len(document_ids) > MAX_DOCUMENTS:
			raise ValueError("Document selection must be a list of up to 10 documents.")
		if any(not isinstance(item, str) for item in document_ids):
			raise ValueError("Document identifiers must be strings.")
		retrieved, sources = retrieve_context(question, list(dict.fromkeys(document_ids)))
	except ValueError as error:
		return jsonify(error=str(error)), 400
	except Exception:
		app.logger.exception("Document retrieval failed")
		return jsonify(error="Document retrieval failed. Check the local embedding model and try again."), 503

	base_url, token = get_colab_config()
	if not base_url or not token:
		return jsonify(error="Colab API is not configured. Add COLAB_API_URL and COLAB_API_TOKEN to .env."), 503

	context_sections = [
		"You are a helpful college study assistant. Be accurate and clear. If the supplied document excerpts do not answer the question, say so rather than guessing.",
	]
	if preferred_name or response_style or learning_goals:
		personalization = []
		if preferred_name:
			personalization.append(f"The student's preferred name is {preferred_name}.")
		if response_style:
			personalization.append(f"Preferred response style: {response_style}.")
		if learning_goals:
			personalization.append(f"Learning goals: {learning_goals}.")
		context_sections.append("Student preferences:\n" + "\n".join(personalization))
	if history:
		context_sections.append(
			"Recent conversation:\n"
			+ "\n".join(f"{item['role']}: {item['content']}" for item in history)
		)
	if retrieved:
		context_sections.append(
			"Relevant excerpts from the student's uploaded documents (use only when relevant):\n"
			+ "\n\n".join(retrieved)
		)

	try:
		response = requests.post(
			f"{base_url}/generate",
			headers={"Authorization": f"Bearer {token}"},
			json={"question": question, "context": "\n\n".join(context_sections)},
			timeout=90,
		)
		response.raise_for_status()
		result = response.json()
		answer = result.get("answer") if isinstance(result, dict) else None
		if not isinstance(answer, str) or not answer.strip():
			raise ValueError("The Colab API returned no answer.")
	except requests.Timeout:
		return jsonify(error="The Colab model took too long to respond. Try again when the runtime is ready."), 504
	except requests.RequestException:
		app.logger.exception("Colab generation request failed")
		return jsonify(error="Could not reach the Colab model. Check that its runtime and tunnel are active."), 502
	except (ValueError, TypeError):
		return jsonify(error="The Colab API returned an invalid response. Check its /generate endpoint."), 502

	return jsonify(answer=answer.strip(), grounded=bool(sources), sources=sources)


@app.get("/api/documents")
def list_documents():
	return jsonify(
		documents=[{"id": item["id"], "name": item["name"]} for item in documents.values()]
	)


@app.post("/api/documents")
def upload_document():
	uploaded = request.files.get("file")
	if uploaded is None or not uploaded.filename:
		return jsonify(error="Choose a PDF, DOCX, or TXT file to upload."), 400
	if len(documents) >= MAX_DOCUMENTS:
		return jsonify(error=f"You can upload up to {MAX_DOCUMENTS} documents at a time."), 413

	safe_name = secure_filename(uploaded.filename)
	extension = safe_name.rsplit(".", 1)[-1].lower() if "." in safe_name else ""
	if extension not in ALLOWED_EXTENSIONS:
		return jsonify(error="Only PDF, DOCX, and TXT files are supported."), 400
	document_id = uuid.uuid4().hex
	path = UPLOAD_DIR / f"{document_id}.{extension}"
	try:
		uploaded.save(path)
		if path.stat().st_size > MAX_UPLOAD_BYTES:
			path.unlink(missing_ok=True)
			return jsonify(error="Files must be 10 MB or smaller."), 413
		text = extract_text(path, extension)
		if not text.strip():
			raise ValueError("No readable text was found in this file.")
		if len(text) > 1_000_000:
			raise ValueError("Extracted text is too large for this demo (maximum 1,000,000 characters).")
		chunks = split_into_chunks(text)
		model = get_embedding_model()
		embeddings = model.encode(chunks, normalize_embeddings=True)
		entry = {
			"id": document_id,
			"name": safe_name or f"document.{extension}",
			"path": path,
			"chunks": chunks,
			"embeddings": embeddings,
		}
		documents[document_id] = entry
	except UnicodeDecodeError:
		path.unlink(missing_ok=True)
		return jsonify(error="TXT files must use UTF-8 text encoding."), 400
	except ValueError as error:
		path.unlink(missing_ok=True)
		return jsonify(error=str(error)), 400
	except Exception:
		path.unlink(missing_ok=True)
		app.logger.exception("Document processing failed")
		return jsonify(error="Could not extract or embed this document. Check the file and local embedding model."), 503

	return jsonify(id=document_id, name=entry["name"]), 201


@app.delete("/api/documents/<document_id>")
def delete_document(document_id):
	entry = documents.pop(document_id, None)
	if entry is None:
		return jsonify(error="Document not found. It may have been removed or the server restarted."), 404
	entry["path"].unlink(missing_ok=True)
	return jsonify(deleted=True)


@app.errorhandler(413)
def request_too_large(_error):
	return jsonify(error="The upload is too large. Files must be 10 MB or smaller."), 413


if __name__ == "__main__":
	app.run(host="127.0.0.1", port=int(os.getenv("PORT", "5000")), debug=False)
